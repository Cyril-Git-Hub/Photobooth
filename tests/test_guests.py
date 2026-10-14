import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

from app.core import guests
from app.core.guests import Guest, GuestListError, load_guests, normalize_text, search_guests


def make_xlsx(path: Path, rows, title="Invités"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    workbook.save(path)


class GuestTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory(prefix="invités é ")
        self.addCleanup(self._tmp.cleanup)
        self.path = Path(self._tmp.name) / "invités.xlsx"

    def test_lecture_standard(self):
        make_xlsx(self.path, [["nom", "prénom", "mail"],
                              ["Martin", "Léa", "LEA@gmail.com "],
                              ["Dupont", "Camille", None]])
        result = load_guests(self.path)
        self.assertEqual(result, [Guest("Dupont", "Camille", ""), Guest("Martin", "Léa", "lea@gmail.com")])
        self.assertEqual(result[1].display_name, "Léa Martin")

    def test_titres_varies_ordre_libre_accents_et_majuscules(self):
        make_xlsx(self.path, [["E-mail", "PRENOM", "Nom de famille", "Remarque"],
                              ["x@y.fr", "Léa", "Martin", "végétarienne"]])
        self.assertEqual(load_guests(self.path), [Guest("Martin", "Léa", "x@y.fr")])

    def test_colonne_mail_absente(self):
        make_xlsx(self.path, [["nom", "prénom"], ["Martin", "Léa"]])
        self.assertEqual(load_guests(self.path), [Guest("Martin", "Léa", "")])

    def test_nom_seul_ou_prenom_seul(self):
        make_xlsx(self.path, [["nom", "prénom", "mail"], ["Martin", None, None], [None, "Zoé", None], [None, None, "orphelin@x.fr"]])
        names = {g.display_name for g in load_guests(self.path)}
        self.assertEqual(names, {"Martin", "Zoé"})                          # la ligne sans nom est ignorée

    def test_lignes_vides_cellules_numeriques_et_espaces(self):
        make_xlsx(self.path, [[None, None, None], ["nom", "prénom", "mail"], [None, None, None],
                              ["  Le   Goff ", "Anne  Sophie", None], [123, 4.5, None]])
        result = load_guests(self.path)
        self.assertIn(Guest("Le Goff", "Anne Sophie", ""), result)           # espaces multiples réduits
        self.assertIn(Guest("123", "4.5", ""), result)

    def test_tri_par_prenom_sans_tenir_compte_des_accents(self):
        make_xlsx(self.path, [["nom", "prénom"], ["A", "Zoé"], ["B", "Émile"], ["C", "Eva"], ["D", "Éric"], ["E", "Adam"]])
        self.assertEqual([g.first_name for g in load_guests(self.path)], ["Adam", "Émile", "Éric", "Eva", "Zoé"])

    def test_fichier_vide_ou_sans_ligne(self):
        make_xlsx(self.path, [])
        self.assertEqual(load_guests(self.path), [])
        make_xlsx(self.path, [["nom", "prénom", "mail"]])
        self.assertEqual(load_guests(self.path), [])

    def test_titres_non_reconnus(self):
        make_xlsx(self.path, [["Martin", "Léa", "lea@x.fr"], ["Dupont", "Camille", None]])
        with self.assertRaises(GuestListError) as ctx:
            load_guests(self.path)
        self.assertIn("nom, prénom, mail", ctx.exception.args[0])

    def test_fichier_absent(self):
        with self.assertRaises(GuestListError) as ctx:
            load_guests(self.path.with_name("absent.xlsx"))
        self.assertIn("introuvable", ctx.exception.args[0])

    def test_fichier_qui_n_est_pas_un_excel(self):
        self.path.write_text("nom;prénom;mail\nMartin;Léa;lea@x.fr\n", encoding="utf-8")
        with self.assertRaises(GuestListError) as ctx:
            load_guests(self.path)
        self.assertIn("Excel", ctx.exception.args[0])

    def test_fichier_ouvert_dans_excel(self):
        make_xlsx(self.path, [["nom", "prénom"], ["Martin", "Léa"]])
        with mock.patch("openpyxl.load_workbook", side_effect=PermissionError("verrouillé")):
            with self.assertRaises(GuestListError) as ctx:
                load_guests(self.path)
        self.assertIn("ouvert dans Excel", ctx.exception.args[0])

    def test_openpyxl_absent(self):
        import sys
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            with self.assertRaises(GuestListError) as ctx:
                load_guests(self.path)
        self.assertIn("pip install", ctx.exception.args[0])


class SearchTests(unittest.TestCase):
    GUESTS = [Guest("Martin", "Léa"), Guest("Dupont", "Jean-Éric"), Guest("Le Goff", "Anne"), Guest("Martin", "Hugo")]

    def test_requete_vide_renvoie_tout(self):
        self.assertEqual(search_guests(self.GUESTS, ""), self.GUESTS)
        self.assertEqual(search_guests(self.GUESTS, "   "), self.GUESTS)

    def test_sans_accents_ni_majuscules(self):
        self.assertEqual(search_guests(self.GUESTS, "lea"), [self.GUESTS[0]])
        self.assertEqual(search_guests(self.GUESTS, "LÉA"), [self.GUESTS[0]])
        self.assertEqual(search_guests(self.GUESTS, "eric"), [self.GUESTS[1]])
        self.assertEqual(search_guests(self.GUESTS, "jean eric"), [self.GUESTS[1]])      # le trait d'union n'empêche rien

    def test_morceau_de_nom_et_prenom(self):
        self.assertEqual({g.first_name for g in search_guests(self.GUESTS, "mar")}, {"Léa", "Hugo"})
        self.assertEqual(search_guests(self.GUESTS, "mar hu"), [self.GUESTS[3]])
        self.assertEqual(search_guests(self.GUESTS, "hu mar"), [self.GUESTS[3]])         # ordre libre
        self.assertEqual(search_guests(self.GUESTS, "le g"), [self.GUESTS[2]])

    def test_aucun_resultat(self):
        self.assertEqual(search_guests(self.GUESTS, "zzz"), [])

    def test_normalize_text(self):
        self.assertEqual(normalize_text("  Jean-Éric  O'Neil "), "jean eric o neil")
        self.assertEqual(normalize_text("Ça Été"), "ca ete")


class TemplateTests(unittest.TestCase):
    def test_modele_cree_puis_lisible(self):
        with tempfile.TemporaryDirectory(prefix="invités é ") as tmp:
            path = Path(tmp) / "sous dossier" / "invites.xlsx"
            self.assertTrue(guests.ensure_template(path))
            self.assertEqual(load_guests(path), [])                                       # titres reconnus, aucun invité
            from openpyxl import load_workbook
            sheet = load_workbook(path).active
            self.assertEqual([c.value for c in sheet[1]], ["nom", "prénom", "mail"])

    def test_fichier_existant_jamais_ecrase(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "invites.xlsx"
            make_xlsx(path, [["nom", "prénom"], ["Martin", "Léa"]])
            self.assertFalse(guests.ensure_template(path))
            self.assertEqual(len(load_guests(path)), 1)


if __name__ == "__main__":
    unittest.main()
